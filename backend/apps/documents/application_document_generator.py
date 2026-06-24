"""
AI Application Document Generator V1 — generate → review → export → save.

Generates STRUCTURED content with Claude (never a final PDF), grounded ONLY in
the user's Smart Profile + application/pack context (no invented facts). The user
reviews, then a deterministic export step renders a real PDF (fpdf2 — selectable
text, never an image) or DOCX (python-docx — editable, ATS-friendly) and stores
it as an encrypted, owner-scoped ``DocumentFile`` (never a raw storage URL).

Gating, AI consent, credit spend (variable per document type), and the budget
guard are enforced by the calling view + ``apps.ai.client.generate`` — never
bypassed here. Export/save make NO AI call and consume NO AI credits.
"""

from __future__ import annotations

import hashlib
import io
import logging

from django.db import transaction

from apps.ai.client import ai_available, generate
from apps.ai.privacy import maybe_redact

from . import document_templates as templates
from .file_encryption import encrypt_bytes_into_record

logger = logging.getLogger(__name__)

FEATURE = "application_document_generation"
_MAX_TOKENS = 3000
_MAX_INSTRUCTIONS = 1500

# --- Structured-output schemas ----------------------------------------------

_QUALITY = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "strengths": {"type": "array", "items": {"type": "string"}},
        "missing_information": {"type": "array", "items": {"type": "string"}},
        "risk_warnings": {"type": "array", "items": {"type": "string"}},
        "suggested_improvements": {"type": "array", "items": {"type": "string"}},
    },
}

_CV_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "title": {"type": "string"},
        "content": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "header": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "name": {"type": "string"},
                        "email": {"type": "string"},
                        "phone": {"type": "string"},
                        "location": {"type": "string"},
                    },
                },
                "summary": {"type": "string"},
                "education": {"type": "array", "items": {"type": "string"}},
                "experience": {"type": "array", "items": {"type": "string"}},
                "projects": {"type": "array", "items": {"type": "string"}},
                "skills": {"type": "array", "items": {"type": "string"}},
                "certifications": {"type": "array", "items": {"type": "string"}},
                "awards": {"type": "array", "items": {"type": "string"}},
                "leadership": {"type": "array", "items": {"type": "string"}},
                "languages": {"type": "array", "items": {"type": "string"}},
            },
        },
        "plain_text": {"type": "string"},
        "quality_checks": _QUALITY,
        "recommended_template": {"type": "string"},
    },
}

_LETTER_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "title": {"type": "string"},
        "content": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "sections": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "additionalProperties": False,
                        "properties": {
                            "heading": {"type": "string"},
                            "body": {"type": "string"},
                        },
                    },
                },
                "closing": {"type": "string"},
            },
        },
        "plain_text": {"type": "string"},
        "quality_checks": _QUALITY,
        "recommended_template": {"type": "string"},
    },
}

_SYSTEM = (
    "You write serious, professional application documents (CVs/résumés, letters, "
    "and emails). You are given ONLY the applicant's own Smart Profile and "
    "application/pack context. CRITICAL: never invent or exaggerate facts — no "
    "fake degrees, employers, dates, skills, awards, or metrics. Use only the "
    "supplied data. When important information is missing, DO NOT fabricate it — "
    "leave it out and add a clear entry to quality_checks.missing_information "
    "instead. Keep the tone professional and avoid exaggerated claims. Tailor the "
    "content to the target application/organization and the requested style. "
    "Return the requested JSON only."
)


# ---- Context ----------------------------------------------------------------


def build_generation_context(user, application=None, bundle=None) -> dict:
    """Deterministic generation context (Smart Profile + application + pack).

    Excludes passport/ID numbers by design (Smart Profile context never includes
    them). Adds requirement titles + the import source URL when available.
    """
    from apps.users.smart_profile import build_application_context_from_profile

    context = build_application_context_from_profile(user, application)

    target_bundle = bundle or (application.linked_bundle if application else None)
    if target_bundle is not None and getattr(target_bundle, "owner_id", None) == user.id:
        from .pack_readiness import build_pack_readiness

        readiness = build_pack_readiness(target_bundle, user)
        context["pack"] = {
            "name": readiness["name"],
            "readiness_score": readiness["score"],
            "missing_documents": [r["title"] for r in readiness["missing_requirements"]],
            "requirements": [
                {"title": r["title"], "description": r.get("description", "")}
                for r in (readiness["required_documents"] + readiness["satisfied_requirements"])
            ][:25],
        }
        # Latest imported requirement source, if any.
        from .models import RequirementExtractionDraft

        draft = (
            RequirementExtractionDraft.objects.filter(owner=user, bundle=target_bundle)
            .exclude(source_url="")
            .order_by("-created_at")
            .first()
        )
        if draft:
            context["source"] = {"url": draft.source_url, "title": draft.page_title}
    return context


def _prompt(context: dict, *, document_type: str, content_style: str,
            target_organization: str, instructions: str) -> str:
    import json

    type_label = templates.DOCUMENT_TYPES.get(
        document_type, templates.DOCUMENT_TYPES["other"]
    )["label"]
    lines = [
        f"DOCUMENT TYPE: {type_label} ({document_type})",
        f"WRITING STYLE: {content_style}",
    ]
    if target_organization:
        lines.append(f"TARGET ORGANIZATION: {target_organization}")
    if instructions:
        lines.append(f"USER INSTRUCTIONS: {instructions[:_MAX_INSTRUCTIONS]}")
    lines.append("\n--- APPLICANT CONTEXT (use ONLY this; do not invent) ---")
    lines.append(json.dumps(context, ensure_ascii=False, default=str))
    return "\n".join(lines)


# ---- Generate ---------------------------------------------------------------


def _clean_str_list(value, cap=40):
    if not isinstance(value, list):
        return []
    out = []
    for v in value[:cap]:
        s = str(v).strip()
        if s:
            out.append(s[:1000])
    return out


def _clean_cv_content(content: dict) -> dict:
    content = content if isinstance(content, dict) else {}
    header = content.get("header") if isinstance(content.get("header"), dict) else {}
    return {
        "header": {
            "name": str(header.get("name", ""))[:200],
            "email": str(header.get("email", ""))[:200],
            "phone": str(header.get("phone", ""))[:60],
            "location": str(header.get("location", ""))[:200],
        },
        "summary": str(content.get("summary", ""))[:2000],
        "education": _clean_str_list(content.get("education")),
        "experience": _clean_str_list(content.get("experience")),
        "projects": _clean_str_list(content.get("projects")),
        "skills": _clean_str_list(content.get("skills")),
        "certifications": _clean_str_list(content.get("certifications")),
        "awards": _clean_str_list(content.get("awards")),
        "leadership": _clean_str_list(content.get("leadership")),
        "languages": _clean_str_list(content.get("languages")),
    }


def _clean_letter_content(content: dict) -> dict:
    content = content if isinstance(content, dict) else {}
    sections = []
    for s in (content.get("sections") or [])[:12]:
        if not isinstance(s, dict):
            continue
        body = str(s.get("body", "")).strip()
        if not body:
            continue
        sections.append({"heading": str(s.get("heading", ""))[:200], "body": body[:4000]})
    return {"sections": sections, "closing": str(content.get("closing", ""))[:600]}


def _clean_quality(q: dict) -> dict:
    q = q if isinstance(q, dict) else {}
    return {
        "strengths": _clean_str_list(q.get("strengths"), 12),
        "missing_information": _clean_str_list(q.get("missing_information"), 12),
        "risk_warnings": _clean_str_list(q.get("risk_warnings"), 12),
        "suggested_improvements": _clean_str_list(q.get("suggested_improvements"), 12),
    }


def generate_application_document(user, payload: dict) -> dict:
    """
    Validate + generate a structured application document (never raises).

    On a genuine model success persists a ``GeneratedApplicationDocument``
    (status=draft) and returns ``{available, reason:"ok", model_called:True,
    generated_document_id, credit_cost, ...}``. Failures (invalid request,
    not-configured, budget, AI error) return ``available:False`` and never persist
    or charge.
    """
    from .models import DocumentBundle, GeneratedApplicationDocument, TrackedApplication

    document_type = payload.get("document_type")
    if document_type not in templates.DOCUMENT_TYPES:
        return {"available": False, "reason": "invalid_request",
                "message": "Unknown document type."}

    if not ai_available():
        return {"available": False, "reason": "not_configured", "model_called": False}

    # Resolve owner-scoped application / bundle (foreign id → blocked).
    application = None
    if payload.get("application_id"):
        application = TrackedApplication.objects.filter(
            pk=payload["application_id"], owner=user
        ).first()
        if application is None:
            return {"available": False, "reason": "not_found",
                    "message": "Application not found."}
    bundle = None
    if payload.get("bundle_id"):
        bundle = DocumentBundle.objects.filter(pk=payload["bundle_id"], owner=user).first()
        if bundle is None:
            return {"available": False, "reason": "not_found",
                    "message": "Pack not found."}

    content_style = templates.normalize_content_style(payload.get("content_style"))
    target_org = str(payload.get("target_organization") or "").strip()[:255]
    if not target_org and application:
        target_org = (application.organization_name or "")[:255]
    instructions = str(payload.get("additional_instructions") or "").strip()

    context = build_generation_context(user, application, bundle)
    prompt = maybe_redact(
        user,
        _prompt(context, document_type=document_type, content_style=content_style,
                target_organization=target_org, instructions=instructions),
    )

    is_cv = templates.is_cv_type(document_type)
    schema = _CV_SCHEMA if is_cv else _LETTER_SCHEMA
    result = generate(
        prompt=prompt, system=_SYSTEM, output_schema=schema,
        max_tokens=_MAX_TOKENS, user=user, feature=FEATURE,
    )
    if not result.ok or not isinstance(result.data, dict):
        reason = "budget" if result.reason == "budget" else "error"
        return {"available": False, "reason": reason, "model_called": False}

    data = result.data
    if is_cv:
        content = _clean_cv_content(data.get("content"))
    else:
        content = _clean_letter_content(data.get("content"))
    quality = _clean_quality(data.get("quality_checks"))
    plain_text = str(data.get("plain_text", "")).strip()[:20000]

    template_key = payload.get("template_key") or templates.recommend_template(
        document_type, getattr(application, "application_type", None), target_org
    )
    if not templates.get_template(template_key):
        template_key = templates.recommend_template(document_type)

    ats_score = None
    warnings = list(quality.get("risk_warnings", []))
    if is_cv:
        ats = templates.validate_ats_structure(content, template_key)
        ats_score = ats["score"]
        warnings = warnings + ats["warnings"]
    warnings = warnings + quality.get("missing_information", [])

    title = str(data.get("title", "")).strip()[:255] or (
        f"{templates.DOCUMENT_TYPES[document_type]['label']}"
        + (f" — {target_org}" if target_org else "")
    )

    gad = GeneratedApplicationDocument.objects.create(
        owner=user,
        application=application,
        bundle=bundle or (application.linked_bundle if application else None),
        document_type=document_type,
        status=GeneratedApplicationDocument.Status.DRAFT,
        title=title,
        target_organization=target_org,
        template_key=template_key,
        content_style=content_style,
        structured_content=content,
        plain_text_preview=plain_text,
        ats_score=ats_score,
        quality_score=None,
        warnings=warnings[:30],
        ai_model=result.model or "",
        credits_charged=0,
    )

    return {
        "available": True,
        "reason": "ok",
        "model_called": True,
        "generated_document_id": gad.id,
        "document_type": document_type,
        "credit_cost": templates.get_document_credit_cost(document_type),
        "title": gad.title,
        "template_key": template_key,
        "recommended_template": templates.recommend_template(document_type),
        "plain_text_preview": plain_text,
        "structured_content": content,
        "quality_checks": quality,
        "ats_score": ats_score,
        "warnings": gad.warnings,
        "available_exports": templates.get_template(template_key)["export_formats"],
    }


# ---- Export (deterministic; NO AI, NO credits) ------------------------------


def _safe_filename(title: str, ext: str) -> str:
    base = "".join(c if (c.isalnum() or c in " -_") else "" for c in title).strip()
    base = " ".join(base.split())[:80] or "document"
    return f"{base}.{ext}"


def export_generated_document(gad, user, *, export_format: str, template_key=None,
                              save_to_pack: bool = False) -> dict:
    """
    Render the reviewed content to a real PDF/DOCX and store it as an encrypted
    DocumentFile (enforcing file + storage limits). No AI call, no credits.
    Optionally attach it to the linked pack. Returns ids + the private download
    route (never a raw storage URL).
    """
    from apps.users import plans

    from .models import DocumentFile, GeneratedApplicationDocument
    from .plan_usage import enforce_plan_limit, enforce_storage_limit

    template_key = template_key or gad.template_key or templates.recommend_template(
        gad.document_type
    )
    if export_format not in ("pdf", "docx"):
        raise ValueError("Unsupported export format.")
    if not templates.template_supports(template_key, gad.document_type, export_format):
        raise ValueError("This template does not support that format for this document.")

    if export_format == "pdf":
        data = _render_pdf(gad, template_key)
        content_type = "application/pdf"
    else:
        data = _render_docx(gad, template_key)
        content_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    filename = _safe_filename(gad.title, export_format)

    # Enforce limits BEFORE persisting (same product limits as any upload).
    enforce_plan_limit(user, plans.RESOURCE_FILES)
    enforce_storage_limit(user, len(data))

    instance = DocumentFile(
        document=None,  # inbox file until/unless saved to the vault/pack
        uploaded_by=user,
        original_filename=filename,
        content_type=content_type,
        file_size=len(data),
        checksum=hashlib.sha256(data).hexdigest(),
    )
    encrypt_bytes_into_record(instance, data, filename)
    instance.save()

    if export_format == "pdf":
        gad.exported_pdf_file = instance
    else:
        gad.exported_docx_file = instance
    gad.template_key = template_key
    if gad.status == GeneratedApplicationDocument.Status.DRAFT:
        gad.status = GeneratedApplicationDocument.Status.EXPORTED
    gad.save(update_fields=[
        "exported_pdf_file", "exported_docx_file", "template_key", "status",
        "updated_at",
    ])

    response = {
        "exported": True,
        "format": export_format,
        "file_id": instance.id,
        # Private, owner-only download route (no raw storage URL).
        "download_url": f"/api/v1/files/{instance.id}/download/",
    }
    if save_to_pack:
        response.update(save_generated_document_to_pack(gad, user, file_instance=instance))
    return response


def save_generated_document_to_pack(gad, user, *, file_instance=None) -> dict:
    """
    Attach an exported file to the user's vault + linked pack as a real Document
    and a satisfied requirement. Enforces the document plan limit. No AI/credits.
    """
    from apps.users import plans

    from .models import (
        Document,
        DocumentBundleRequirement,
        GeneratedApplicationDocument,
    )
    from .plan_usage import enforce_plan_limit

    file_instance = file_instance or gad.exported_pdf_file or gad.exported_docx_file
    if file_instance is None:
        return {"saved_to_pack": False, "reason": "no_export"}

    with transaction.atomic():
        if gad.created_document is None:
            enforce_plan_limit(user, plans.RESOURCE_DOCUMENTS)
            doc = Document.objects.create(
                owner=user,
                title=gad.title[:255],
                document_type=templates.DOCUMENT_TYPES[gad.document_type]["label"][:100],
                notes="Generated by CertaNest — review before submitting.",
            )
            gad.created_document = doc
        else:
            doc = gad.created_document
        # Attach the exported file to the vault document.
        if file_instance.document_id is None:
            file_instance.document = doc
            file_instance.save(update_fields=["document"])

        bundle = gad.bundle
        requirement_id = None
        if bundle is not None and bundle.owner_id == user.id:
            req = DocumentBundleRequirement.objects.create(
                owner=user,
                bundle=bundle,
                title=gad.title[:255],
                requirement_type=DocumentBundleRequirement.RequirementType.DOCUMENT,
                is_required=False,
                status=DocumentBundleRequirement.Status.ATTACHED,
                linked_document=doc,
                linked_file=file_instance,
            )
            requirement_id = req.id
            bundle.recalculate_readiness()

        gad.status = GeneratedApplicationDocument.Status.SAVED_TO_PACK
        gad.save(update_fields=["created_document", "status", "updated_at"])

    out = {"saved_to_pack": True, "document_id": doc.id, "requirement_id": requirement_id}
    if gad.bundle is not None:
        from .pack_readiness import build_pack_readiness

        out["pack_readiness"] = build_pack_readiness(gad.bundle, user)
    return out


# ---- Renderers (fpdf2 / python-docx) ---------------------------------------


def _cv_sections(content: dict):
    order = [
        ("Summary", "summary"), ("Education", "education"), ("Experience", "experience"),
        ("Projects", "projects"), ("Skills", "skills"),
        ("Certifications", "certifications"), ("Awards", "awards"),
        ("Leadership", "leadership"), ("Languages", "languages"),
    ]
    for label, key in order:
        value = content.get(key)
        if not value:
            continue
        yield label, value


def _render_pdf(gad, template_key: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()
    content = gad.structured_content or {}

    def _cell(text, h):
        # new_x/new_y keep the cursor at the left margin on the next line so the
        # following multi_cell always has the full printable width (fpdf2 2.8).
        pdf.multi_cell(0, h, _latin(text), new_x="LMARGIN", new_y="NEXT")

    def heading(text):
        pdf.set_font("Helvetica", "B", 12)
        _cell(text, 6)
        pdf.ln(1)

    def body(text, bullet=False):
        pdf.set_font("Helvetica", size=10.5)
        prefix = "-  " if bullet else ""
        _cell(prefix + text, 5)

    if templates.is_cv_type(gad.document_type):
        header = content.get("header") or {}
        pdf.set_font("Helvetica", "B", 18)
        _cell(header.get("name") or gad.title, 9)
        contact = " | ".join(
            x for x in [header.get("email"), header.get("phone"), header.get("location")] if x
        )
        if contact:
            pdf.set_font("Helvetica", size=10)
            _cell(contact, 5)
        pdf.ln(2)
        for label, value in _cv_sections(content):
            heading(label)
            if isinstance(value, list):
                for item in value:
                    body(str(item), bullet=True)
            else:
                body(str(value))
            pdf.ln(2)
    else:
        pdf.set_font("Helvetica", "B", 15)
        _cell(gad.title, 8)
        pdf.ln(3)
        for section in (content.get("sections") or []):
            if section.get("heading"):
                heading(section["heading"])
            body(section.get("body", ""))
            pdf.ln(2)
        if content.get("closing"):
            pdf.ln(2)
            body(content["closing"])
        if not content.get("sections") and gad.plain_text_preview:
            body(gad.plain_text_preview)

    return bytes(pdf.output())


def _render_docx(gad, template_key: str) -> bytes:
    from docx import Document as Docx
    from docx.shared import Pt

    doc = Docx()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    content = gad.structured_content or {}

    if templates.is_cv_type(gad.document_type):
        header = content.get("header") or {}
        doc.add_heading(header.get("name") or gad.title, level=0)
        contact = " | ".join(
            x for x in [header.get("email"), header.get("phone"), header.get("location")] if x
        )
        if contact:
            doc.add_paragraph(contact)
        for label, value in _cv_sections(content):
            doc.add_heading(label, level=1)
            if isinstance(value, list):
                for item in value:
                    doc.add_paragraph(str(item), style="List Bullet")
            else:
                doc.add_paragraph(str(value))
    else:
        doc.add_heading(gad.title, level=0)
        for section in (content.get("sections") or []):
            if section.get("heading"):
                doc.add_heading(section["heading"], level=1)
            doc.add_paragraph(section.get("body", ""))
        if content.get("closing"):
            doc.add_paragraph(content["closing"])
        if not content.get("sections") and gad.plain_text_preview:
            doc.add_paragraph(gad.plain_text_preview)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _latin(text: str) -> str:
    """fpdf2 core fonts are latin-1; replace unsupported glyphs safely."""
    return str(text).encode("latin-1", "replace").decode("latin-1")
